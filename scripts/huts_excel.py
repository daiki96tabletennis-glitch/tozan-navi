#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""山小屋の情報と、自動更新の有無を、管理表（yamatch_kensho_kanri (3).xlsx）の「山小屋情報」シートに書き出す。

  python3 scripts/huts_excel.py

data/huts.json などから毎回まるごと作り直す（シートを手で直さない）。山小屋のデータが変わったら実行する
（定期更新の分を GitHub から取り込んだあと・山小屋の情報を直したあと）。
表示の値（予約の開始・営業期間・定員）は gen_huts_page.py と同じ順で選ぶ。
"""
import json, os, re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'yamatch_kensho_kanri (3).xlsx')
SHEET = '山小屋情報'
NOTE_SHEET = '山小屋情報の見方'
BASE = 'https://tozan-navi.com'


def load(name):
    p = os.path.join(ROOT, 'data', name)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else {}


def shown_start(h):
    """gen_huts_page.py と同じ順で、ページに出る「予約の開始」と、その出どころを返す"""
    au = h.get('auto') or {}
    stale = h.get('bookingStale')
    if (au.get('bookingStart') or {}).get('value'):
        return au['bookingStart']['value'], '自動（公式ページのレシピ）'
    start = None if stale else h.get('bookingStart')
    src = '手入力（公式サイト・出典を読んで確認）' if start else ''
    for key, label in (('bookingStartDate', '自動（公式ページの日付）'), ('bookingStartWeb', '自動（やまたんのお知らせ）')):
        web = au.get(key) or {}
        if web:
            yrs = [int(x) for x in re.findall(r'(20\d\d)年', start or '')]
            if not start or (yrs and max(yrs) < (web.get('year') or 0)):
                return web['value'], label
    if not start and (au.get('bookingWindow') or {}).get('value'):
        return au['bookingWindow']['value'], '自動（やまたんの何日前）'
    if h.get('bookingNone') and not stale:
        return '予約不要', '手入力'
    if stale:
        return '確認中', '出典の記載が変わった'
    if h.get('bookingStartUnstated'):
        return '記載なし', '手入力（公式サイトに記載が無いと確認）'
    return start or '', src


def main():
    H = load('huts.json')['huts']
    M = {m['id']: m for m in json.load(open(os.path.join(ROOT, 'data', 'mountains.json'), encoding='utf-8'))}
    src = load('huts_sources.json')
    yamatan = (src.get('huts') or {})
    news_names = set(((src.get('yamatanNews') or {}).get('names') or {}).values())
    recipes = {}
    for r in (load('huts_recipes.json').get('recipes') or []):
        recipes.setdefault(r['hut'], []).append(r['field'])
    sns = (load('huts_sns.json').get('huts') or {})
    rules = (load('huts_rules.json').get('huts') or {})
    RT = {'rel': '何日前・何か月前', 'mon': '月ごと', 'fix': '一斉（決まった日）', 'phase': '期間区切り'}

    head = ['ID', '山小屋', '地域', '対応する山', '別ルート・縦走で使う山', '公式サイト', '電話',
            '営業期間（ページの表示）', '営業期間の出どころ', '公式サイトの注記',
            '予約の開始（ページの表示）', '予約の開始の出どころ', '予約方法', '予約必須', '定員（ページの表示）',
            'お知らせの型', 'ページの状態', '検索に出す', '出典（一覧）', '出典の更新日', '最終確認日',
            '自動で更新：営業期間・予約の区分', '自動取得：やまたん', '自動取得：公式ページのレシピ', '自動取得：型で読む',
            '自動取得：Instagram', '予約開始日の監視', '定期の見張り', '手で書いた値の扱い', 'ページ']
    rows = []
    for h in H:
        au = h.get('auto') or {}
        start, start_src = shown_start(h)
        sf = h.get('seasonFix') or {}
        season = sf.get('value') or (au.get('season') or {}).get('value')
        season_src = ('手入力（公式・自治体の案内で確認 %s）' % sf.get('checkedAt')) if sf.get('value') else ('自動（公式・予約サイト）' if season else '出典の記載')
        season = season or re.sub(r'^2026年の営業：', '', h.get('openText') or '') or ''
        cap = (au.get('capacity') or {}).get('value') or h.get('capacity')
        ref = h.get('capacityRef')
        cf = h.get('capacityFix')
        cap_txt = (cf.get('text') or '%s人' % cf['value']) + '（公式・自治体の案内で確認）' if cf else ('%s人' % cap) if cap else (('約%s人（山と溪谷オンライン）' % ref) if isinstance(ref, int) else (ref or ''))
        su = h.get('sourceUrl') or ''
        auto_season = '自動（毎週。出典の原文）' if ('pref.nagano' in su or 'yamakei-online.com/yama-ya' in su or 'lodge/list' in su) else '手入力（自動では変わらない）'
        yt = (yamatan.get(h['name']) or {}).get('yamatan')
        yt_txt = ''
        if yt:
            yt_txt = 'あり（定員・営業期間・何日前から）'
        if h['name'] in news_names:
            yt_txt = (yt_txt + '／' if yt_txt else '') + '予約開始のお知らせ'
        r = rules.get(h['id']) or {}
        rt = (r.get('r') or {}).get('t')
        if rt in ('rel', 'mon'):
            watch = '決まりが変わっていないか（12〜6月は週1回、それ以外は毎月1日）'
        elif rt in ('fix', 'phase'):
            watch = '来季の予約開始日（発表が出そうな期間は3日に1回、それ以外は毎月1日）'
        else:
            watch = '毎月1日に公式ページを見る（変わると「変更あり」）' if h.get('officialUrl') else ''
        # gen_huts_page.py と同じ：予約不要でもなく、予約の開始（「確認中」「記載なし」は除く）も方法も無い小屋は検索に出さない
        # 予約開始日の監視（huts_auto.py と同じ条件）：自動でページに反映／報告だけ／監視なし
        rec = set(recipes.get(h['name'], []))
        mon = []
        if h.get('bookingNone'):
            mon = ['対象外（予約不要）']
        else:
            if rec & {'bookingStart', 'bookingStartDate'}:
                mon.append('自動で反映：公式ページのレシピ')
            if h['name'] in news_names:
                mon.append('自動で反映：やまたんの予約開始のお知らせ（手の値より新しい年のとき）')
            if rt == 'fix' and h.get('officialUrl') and not (rec & {'bookingStart', 'bookingStartDate'}):
                mon.append('自動で反映：型で読む（年が今年か来年・前の年の日付から45日以内・候補が1つのときだけ。それ以外は報告）')
            if rt in ('rel', 'mon'):
                mon.append('報告だけ：「◯か月前」の数字の変化と、一斉解放の書き方が出てきたか')
            if rt == 'phase':
                mon.append('報告だけ：期間区切りは自動では替えない（公式の記載が変わると「変更あり」）')
            if not rt and h.get('officialUrl'):
                mon.append('報告だけ：予約の開始らしい記載が見つかったら候補として報告（未登録の小屋）')
            if not mon:
                mon = ['監視なし（公式サイトのURLが無い）']
        thin = not (h.get('bookingNone') or start not in ('', '確認中', '記載なし') or h.get('bookingMethods'))
        state = '確認中' if h.get('bookingStale') else ('変更あり' if h.get('officialChanged') else '通常')
        ig = (sns.get(h['id']) or {}).get('instagram') or []
        rows.append([
            h['id'], h['name'], h['region'],
            '・'.join(M[i]['name'] for i in h['mountainIds'] if i in M),
            '・'.join('%s（%s）' % (M[o['id']]['name'], o.get('relation') or '') for o in h.get('otherMountains') or [] if o['id'] in M),
            (h.get('officialUrl') or '').strip(), h.get('tel') or '',
            season, season_src, '／'.join(x for x in (h.get('officialNote'), sf.get('note')) if x),
            start, start_src, '・'.join(h.get('bookingMethods') or []) if not h.get('bookingNone') else '予約なし',
            'はい' if h.get('bookingRequired') else '', cap_txt,
            RT.get(rt, '非対応'), state, 'いいえ（noindex）' if thin else 'はい',
            su, h.get('sourceUpdated') or '', h.get('lastVerified') or '',
            auto_season, yt_txt, '・'.join(sorted(set(recipes.get(h['name'], [])))),
            'あり（公式サイトの予約・お知らせのページ）' if h.get('officialUrl') else 'なし（公式サイトのURLなし）',
            ('準備中（鍵が未設定）：' + '・'.join('@' + x for x in ig[:4]) + ('ほか' if len(ig) > 4 else '')) if ig else '',
            '／'.join(mon), watch, '書き換えない（予約の開始・方法・定員は人が確かめて直す）。自動で読めた値が確認を通れば、ページではそちらを優先して表示する',
            BASE + '/huts/%s/' % h['id']])

    wb = openpyxl.load_workbook(XLSX)
    for name in (SHEET, NOTE_SHEET):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet(SHEET)
    ws.append(head)
    for r in rows:
        ws.append(r)
    fill = PatternFill('solid', fgColor='4E6535')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    widths = [14, 22, 12, 16, 22, 30, 14, 28, 18, 24, 28, 24, 20, 8, 16, 14, 10, 12, 30, 12, 12, 22, 26, 22, 24, 26, 44, 32, 30, 36]
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    ns = wb.create_sheet(NOTE_SHEET)
    for line in [
        ['山小屋情報シートの見方（scripts/huts_excel.py が data/huts.json などから毎回作り直す。手で直さない）'],
        [''],
        ['ページの表示', '予約の開始・営業期間・定員は、山小屋ページ（/huts/）に出ている値。選ぶ順は gen_huts_page.py と同じ'],
        ['予約の開始の出どころ', '手入力＝公式サイト・出典を読んで書き写した値（huts_booking_extra.py など）。自動＝公式ページのレシピ・決まった書き方の日付・やまたん'],
        ['自動で更新：営業期間・予約の区分', '毎週の定期更新（huts-update.yml）が出典（長野県の山小屋ポータル・山と溪谷オンラインの山小屋リスト）を取り直して書き換える'],
        ['自動取得：やまたん', '予約サイト「やまたん」の小屋ページ（data/huts_sources.json）から定員・営業期間・何日前から予約できるかを読む。「予約開始のお知らせ」は年度ごとのページ'],
        ['自動取得：公式ページのレシピ', '公式ページの決まった言い回しを読む正規表現（data/huts_recipes.json）。bookingStart＝予約の開始の文、bookingStartDate＝予約開始日、season＝営業期間'],
        ['自動取得：型で読む', '小屋ごとの設定なしで、公式サイトの予約・お知らせのページから予約開始の書き方の型を読む（huts_patterns.py）。自動で使うのは条件を満たす一斉型だけ。ほかは報告に出すだけ'],
        ['自動取得：Instagram', 'Instagram の投稿を読む仕組み（huts_instagram.py）。鍵（IG_USER_ID / IG_TOKEN）が未設定のあいだは動かない'],
        ['定期の見張り', 'update_huts.py --scheduled が見に行く頻度'],
        ['予約開始日の監視', '自動で反映＝読み取った日付が確認（年が今年か来年・実在する日付・前の年から45日以内・候補が1つ）を通ると、ページの予約の開始と、予約開始のお知らせの日付が替わる（data/huts_history.json に年ごとに残る）。報告だけ＝管理画面のお知らせに出るので、人が公式サイトで確かめて直す'],
        ['手で書いた値の扱い', '予約の開始・予約方法・定員の手で書いた値（huts_booking_extra.py など）は自動では書き換えない。出典や公式ページの記載が変わると、ページで「確認中」「変更あり」になり、管理画面のお知らせに出る'],
        ['ページの状態', '通常／確認中（出典の記載が変わった）／変更あり（公式サイトの記載が変わった）'],
        ['検索に出す', '予約の開始も方法も分からない小屋のページは noindex（検索に出さない・サイトマップに入れない）'],
    ]:
        ns.append(line)
    ns.column_dimensions['A'].width = 34
    ns.column_dimensions['B'].width = 120
    ns['A1'].font = Font(bold=True)
    wb.save(XLSX)
    print('%s：%d軒を書き出しました（%s）' % (SHEET, len(rows), os.path.basename(XLSX)))


if __name__ == '__main__':
    main()
